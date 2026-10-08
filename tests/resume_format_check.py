# -*- coding: utf-8 -*-
r"""
简历多格式上传测试脚本（2026-10-07）
====================================
验证 /api/resume/upload 对 PDF / DOCX / TXT 三种格式的支持，
以及非法文件的拦截（老版 .doc、伪 docx 改名文件）。

测试矩阵：
  [1] DOCX 简历（含段落+表格）     → 期望 200 出题成功
  [2] TXT 简历                     → 期望 200 出题成功
  [3] 老版 .doc                    → 期望 400，提示另存为 .docx/PDF
  [4] 伪 DOCX（PK 头 + 垃圾内容）  → 期望 422，提示解析失败
  [5] 真实 PDF（张三丰简历）       → 期望 200 出题成功（MinerU 链路，耗时较长）

用法：app\venv\Scripts\python.exe tests\resume_format_check.py
注意：场景 [1][2][5] 会真实创建面试记录并扣减配额、调用 AI 出题。
"""
import json
import os
import sys

import requests

BASE = "http://127.0.0.1:10009"
USERNAME = "ui_tester_2026"
PASSWORD = "UiTester2026"

DOCX_PATH = "app/temp/test_resumes/resume_test.docx"
TXT_PATH = "app/temp/test_resumes/resume_test.txt"
PDF_PATH = "docs/resumes/张三丰前端开发工程师简历.pdf"


def ensure_test_files():
    """生成 DOCX/TXT 测试简历（自包含：复跑脚本时自动重建，PDF 用项目自带测试简历）"""
    os.makedirs("app/temp/test_resumes", exist_ok=True)

    # TXT 简历（后端开发方向）
    if not os.path.exists(TXT_PATH):
        txt = (
            "姓名：陈明\n电话：13800138002\n邮箱：chenming@test.com\n"
            "教育背景\n2020-09 至 2024-06 华中科技大学\n专业：计算机科学与技术 学位：本科\n"
            "工作经历\n2024-07 至今 字节跳动公司\n职位：后端开发工程师\n"
            "职责：负责推荐服务后端开发，使用 Python 和 Go，优化接口性能\n"
            "项目经验\n推荐服务重构\n描述：重构老推荐服务，QPS 提升 3 倍\n"
            "技术栈：Python, FastAPI, Redis, MySQL\n技能\nPython, Go, FastAPI, Redis, MySQL, Docker\n证书\nCET-6\n"
        )
        with open(TXT_PATH, "w", encoding="utf-8") as f:
            f.write(txt)

    # DOCX 简历（前端开发方向，含段落 + 表格）
    if not os.path.exists(DOCX_PATH):
        from docx import Document
        doc = Document()
        doc.add_paragraph("姓名：刘婷")
        doc.add_paragraph("电话：13900139003")
        doc.add_paragraph("邮箱：liuting@test.com")
        doc.add_paragraph("教育背景")
        t1 = doc.add_table(rows=1, cols=3)
        t1.rows[0].cells[0].text = "2019-09 至 2023-06 武汉大学"
        t1.rows[0].cells[1].text = "专业：软件工程"
        t1.rows[0].cells[2].text = "学位：本科"
        doc.add_paragraph("工作经历")
        doc.add_paragraph("2023-07 至今 腾讯公司")
        doc.add_paragraph("职位：前端开发工程师")
        doc.add_paragraph("职责：负责小程序前端开发，使用 Vue 和 TypeScript，优化页面加载速度")
        doc.add_paragraph("项目经验")
        doc.add_paragraph("企业微信管理后台")
        doc.add_paragraph("描述：重写管理后台前端，首屏时间缩短 60%")
        doc.add_paragraph("技术栈：Vue3, TypeScript, Vite, Pinia")
        doc.add_paragraph("技能")
        doc.add_paragraph("JavaScript, TypeScript, Vue3, Vite, Node.js, Git")
        doc.save(DOCX_PATH)


def login():
    """登录测试账号，返回 Bearer token"""
    resp = requests.post(
        f"{BASE}/api/auth/login",
        json={"username": USERNAME, "password": PASSWORD},
        timeout=10,
    )
    data = resp.json()
    if resp.status_code != 200 or not data.get("auth_token"):
        print(f"[登录失败] {resp.status_code} {data}")
        sys.exit(1)
    return data["auth_token"]


def get_quota(token):
    """查询当前剩余面试次数"""
    resp = requests.get(
        f"{BASE}/api/user/quota",
        headers={"Authorization": f"Bearer {token}"},
        timeout=10,
    )
    if resp.status_code == 200:
        data = resp.json()
        # 配额字段名因接口实现而异，把整个返回打出来兜底
        return data
    return {"status_code": resp.status_code}


def upload(token, path, position, filename=None):
    """上传简历文件，返回 (状态码, 响应JSON)"""
    with open(path, "rb") as f:
        files = {"resume": (filename or path.split("/")[-1].split("\\")[-1], f)}
        resp = requests.post(
            f"{BASE}/api/resume/upload",
            headers={"Authorization": f"Bearer {token}"},
            files=files,
            data={"work_experience": "1", "target_position": position},
            timeout=300,  # PDF 走 MinerU 可能要几十秒
        )
    try:
        return resp.status_code, resp.json()
    except Exception:
        return resp.status_code, {"raw": resp.text[:200]}


def check(label, ok, detail):
    mark = "PASS" if ok else "FAIL"
    print(f"[{mark}] {label}\n       {detail}\n")


def main():
    print("=" * 60)
    print("简历多格式上传测试")
    print("=" * 60)
    ensure_test_files()
    token = login()
    print(f"登录成功，剩余配额信息: {get_quota(token)}\n")

    # [1] DOCX：段落 + 表格混合简历
    code, data = upload(token, DOCX_PATH, "前端开发")
    check("场景[1] DOCX 简历上传", code == 200 and data.get("status") == "success",
          f"HTTP {code} → {json.dumps(data, ensure_ascii=False)[:160]}")

    # [2] TXT：纯文本简历
    code, data = upload(token, TXT_PATH, "后端开发")
    check("场景[2] TXT 简历上传", code == 200 and data.get("status") == "success",
          f"HTTP {code} → {json.dumps(data, ensure_ascii=False)[:160]}")

    # [3] 老版 .doc：伪造一个（内容随意，扩展名命中 doc 分支即被拦截）
    doc_path = "app/temp/test_resumes/fake_old.doc"
    with open(doc_path, "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0" + b"fake old doc content" * 4)
    code, data = upload(token, doc_path, "前端开发")
    msg = data.get("message") or data.get("detail") or ""
    check("场景[3] 老版 .doc 拦截", code == 400 and "另存为" in msg,
          f"HTTP {code} → {msg}")

    # [4] 伪 DOCX：PK 头但不是合法 docx（改名攻击），应走到解析环节并失败
    fake_path = "app/temp/test_resumes/fake_resume.docx"
    with open(fake_path, "wb") as f:
        f.write(b"PK\x03\x04" + b"not a real docx" * 16)
    code, data = upload(token, fake_path, "前端开发")
    msg = data.get("message") or data.get("detail") or ""
    err_code = data.get("error_code") or data.get("detail", "")
    ok = code == 422 and "解析失败" in msg
    check("场景[4] 伪 DOCX（PK头改名文件）拦截", ok,
          f"HTTP {code} → {msg} (error_code={err_code})")

    # [5] 真实 PDF：MinerU 主路径（GPU 解析，耗时 30~120 秒属正常）
    code, data = upload(token, PDF_PATH, "前端开发")
    check("场景[5] 真实 PDF 上传（MinerU 链路）",
          code == 200 and data.get("status") == "success",
          f"HTTP {code} → {json.dumps(data, ensure_ascii=False)[:160]}")

    print(f"测试后配额信息: {get_quota(token)}")
    print("测试结束")


if __name__ == "__main__":
    main()
